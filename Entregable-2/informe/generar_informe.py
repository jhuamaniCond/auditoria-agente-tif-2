# -*- coding: utf-8 -*-
"""
Genera Informe_Entregable_2.docx del TIF (Grupo 6 - Agente de Evidencia y
Auditoria). El .docx es una SALIDA: cualquier cambio se hace aqui y se vuelve
a ejecutar.

    pip install python-docx
    python Entregable-2/informe/generar_informe.py

Reutiliza (sin modificarlo) el motor de estilo del Entregable 1
(Entregable-1/informe/docx_builder.py): Times New Roman 12, tablas simples y
rojo solo para lo pendiente.

La seccion "Resultados sobre la UNSA" se arma en tiempo de ejecucion con:
  - Entregable-2/datos/documentos_unsa/fuentes.json (documentos publicos y URLs)
  - la matriz de cobertura del proyecto 'unsa' (agente/herramientas.py), si se
    puede importar; si no, deja un marcador en rojo.
"""

import contextlib
import datetime as dt
import io
import json
import os
import sys

AQUI = os.path.dirname(os.path.abspath(__file__))
ENTREGABLE2 = os.path.dirname(AQUI)
REPO = os.path.dirname(ENTREGABLE2)
AGENTE = os.path.join(ENTREGABLE2, "agente")
DATOS = os.path.join(ENTREGABLE2, "datos")
FUENTES = os.path.join(DATOS, "documentos_unsa", "fuentes.json")
SALIDA = os.path.join(AQUI, "Informe_Entregable_2.docx")

try:
    import docx  # noqa: F401
except ImportError:
    sys.exit("Falta python-docx: pip install python-docx")

sys.path.insert(0, os.path.join(REPO, "Entregable-1", "informe"))
from docx_builder import Documento, _escribir  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

INTEGRANTES = [
    "Huamani Condori, Jeanpiero Sixto",
    "Mendoza Contreras, Giovani Angel",
    "Suarez Huamani, Marco Antonio",
    "Yunque Quispe, Gustavo Alonso",
]
DOCENTE = "Edward Vicente Zarate Carlos"


# --------------------------------------------------------------------------
# utilidades
def codigo(doc, lineas):
    """Bloque monoespaciado (diagramas y comandos), una linea por parrafo."""
    for linea in lineas:
        p = doc.parrafo("`%s`" % linea if linea.strip() else " ", espacio_despues=0,
                        sangria_izq=0.5)
        p.alignment = 0
    doc.espacio(6)


def indice(doc):
    """Campo TOC de Word; se rellena al abrir el documento (updateFields)."""
    doc.parrafo("**Índice**", centrado=True)
    p = doc.doc.add_paragraph()

    def fld(tipo):
        r = p.add_run()
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), tipo)
        r._r.append(el)

    fld("begin")
    r = p.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-2" \\h \\z \\u'
    r._r.append(instr)
    fld("separate")
    _escribir(p, "[Índice: si no aparece, clic derecho aquí y «Actualizar campo» (F9).]",
              rojo=True)
    fld("end")
    ajustes = doc.doc.settings.element
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    ajustes.append(upd)
    doc.salto_pagina()


def leer_fuentes():
    """Lista de documentos publicos de la UNSA (estructura tolerante)."""
    try:
        with open(FUENTES, encoding="utf-8") as f:
            datos = json.load(f)
    except (OSError, ValueError):
        return None
    if isinstance(datos, dict):
        for clave in ("documentos", "fuentes", "items", "archivos"):
            if isinstance(datos.get(clave), list):
                return datos[clave]
        return [dict(v, archivo=k) if isinstance(v, dict) else {"archivo": k, "url": v}
                for k, v in datos.items()]
    return datos if isinstance(datos, list) else None


def _campo(d, *claves):
    for c in claves:
        v = d.get(c)
        if v:
            return ", ".join(v) if isinstance(v, list) else str(v)
    return "-"


def leer_matriz():
    """Matriz de cobertura y estado del expediente 'unsa' (o None si no se puede)."""
    sys.path.insert(0, AGENTE)
    try:
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            import herramientas
            cob = herramientas.matriz_cobertura("unsa")
            est = herramientas.estado_expediente("unsa")
        return cob, est
    except Exception as e:  # noqa: BLE001  (el informe no debe abortar)
        return None, str(e)


def contar_lineas(nombre):
    try:
        with open(os.path.join(DATOS, nombre), encoding="utf-8") as f:
            return sum(1 for linea in f if linea.strip())
    except OSError:
        return 0


# --------------------------------------------------------------------------
def caratula(doc):
    doc.titulo_portada("UNIVERSIDAD NACIONAL DE SAN AGUSTÍN DE AREQUIPA")
    doc.titulo_portada("FACULTAD DE INGENIERÍA DE PRODUCCIÓN Y SERVICIOS")
    doc.titulo_portada("ESCUELA PROFESIONAL DE INGENIERÍA DE SISTEMAS")
    doc.espacio(40)
    doc.titulo_portada("AUDITORÍA DE SISTEMAS")
    doc.titulo_portada("Trabajo de Investigación Formativa (TIF) — Segundo entregable")
    doc.espacio(20)
    doc.titulo_portada("Proyecto agéntico: Agente de Evidencia y Auditoría para el PIMS "
                       "multiagente basado en ISO/IEC 27701:2025, aplicado a la UNSA")
    doc.espacio(20)
    doc.titulo_portada("Grupo 6 — Agente de Evidencia y Auditoría", negrita=False)
    doc.espacio(30)
    doc.parrafo("**Integrantes:**", centrado=True)
    for nombre in INTEGRANTES:
        doc.parrafo(nombre, centrado=True, espacio_despues=0)
    doc.espacio(20)
    p = doc.parrafo("**Docente:** ", centrado=True)
    _escribir(p, DOCENTE)
    doc.espacio(40)
    doc.parrafo("Octubre de 2026", centrado=True)
    doc.salto_pagina()


def introduccion(doc):
    doc.titulo(1, "1. Introducción")
    doc.parrafo(
        "En el primer entregable se transformó la tarea de preparar una auditoría del PIMS en "
        "la skill `evidencia-auditoria-pims` y en el asistente local AUDITA-PIMS: registro de "
        "evidencias con SHA-256 y cadena de custodia, verificación de integridad, matriz de "
        "cobertura contra ISO/IEC 27701:2025 y hallazgos preliminares redactados por un modelo "
        "local con RAG. Todos los pasos los iniciaba una persona desde la aplicación.")
    doc.parrafo(
        "Este segundo entregable convierte esa skill en un **agente**: un sistema con objetivo, "
        "datos, contexto y herramientas definidos, que encadena los pasos por sí mismo en "
        "flujos orquestados, se coordina con otros agentes del PIMS multiagente (en particular "
        "con el Agente de Incidentes), expone sus capacidades por el Model Context Protocol "
        "(MCP) y se detiene en puertas de aprobación humana antes de cualquier conclusión. Cubre "
        "las tres tareas restantes del TIF: diseño de agente, orquestación y automatización, y "
        "proyecto agéntico.")
    doc.parrafo(
        "Para que el resultado sea una aplicación real y no un ejemplo ficticio, el agente se "
        "aplicó a la **Universidad Nacional de San Agustín de Arequipa (UNSA)**, usando solo "
        "documentos públicos descargados de sus sitios oficiales. Lo que la universidad no "
        "publica no se sustituye con documentos inventados: queda en estado *Sin evidencia* y "
        "se reporta como hallazgo preliminar.")
    doc.titulo(2, "1.1. Objetivos")
    doc.parrafo(
        "**Objetivo general.** Construir un agente local de evidencia y auditoría que mantenga "
        "íntegro y trazable el expediente del PIMS y prepare la auditoría interna (cláusula 9.2) "
        "de forma repetible, coordinado con otros agentes y con validación humana obligatoria.")
    doc.parrafo("**Objetivos específicos:**")
    doc.numerada([
        "Diseñar el agente definiendo su objetivo SMART, alcance, datos, contexto, herramientas, "
        "memoria, políticas, puertas humanas, métricas y riesgos.",
        "Publicar la skill del primer entregable como una capa de herramientas con contrato "
        "JSON, reutilizable desde Python, MCP y HTTP.",
        "Orquestar los flujos `auditoria_interna`, `incidente` y `bandeja` con LangGraph, con "
        "estado persistido y puertas de aprobación del auditor y del DPO.",
        "Coordinar el agente con el resto del PIMS mediante un bus de eventos y un servidor MCP.",
        "Automatizar el ingreso de evidencias y la verificación periódica de integridad.",
        "Aplicar el proyecto a la UNSA con documentos públicos reales y validarlo con casos de "
        "prueba.",
    ])


def diseno(doc):
    doc.titulo(1, "2. Diseño del agente")
    doc.parrafo(
        "El agente es el Agente 4 del PIMS multiagente propuesto en la guía del curso. Su "
        "ficha completa está en `01-diseno-agente/diseno_agente.md`; aquí se resumen sus "
        "decisiones de diseño.")
    doc.titulo(2, "2.1. Objetivo y alcance")
    doc.tabla([
        ["Criterio", "Definición"],
        ["Específico", "Para el proyecto `unsa`: verificar la integridad de todas las evidencias, "
         "calcular la matriz de los 38 controles aplicables al rol responsable, proponer "
         "pertinencia y hallazgos, y emitir informe y plan de acción solo tras aprobación del "
         "auditor."],
        ["Medible", "100 % de evidencias con hash verificado; 100 % de controles no cubiertos con "
         "hallazgo; 0 controles cubiertos sin evidencia; 100 % de decisiones y llamadas MCP en "
         "bitácora."],
        ["Alcanzable", "Reutiliza el motor determinista de la skill y el RAG del Entregable 1; el "
         "LLM solo clasifica, evalúa pertinencia y redacta."],
        ["Relevante", "Responde a 9.1, 9.2 y A.3.14 y es la fuente de evidencias del resto de "
         "agentes."],
        ["Temporal", "Auditoría en modo rápido en menos de 2 minutos en la laptop del equipo "
         "(i5-6200U, sin GPU); entrega funcional al cierre del semestre 2026-II."],
    ], anchos_cm=[3.0, 13.0])
    doc.parrafo(
        "**Dentro del alcance:** registrar y verificar evidencias, calcular cobertura, proponer "
        "pertinencia y hallazgos, documentar como evidencia los reportes de otros agentes y "
        "emitir el informe tras aprobación. **Fuera del alcance:** decidir el cumplimiento o "
        "clasificar no conformidades (auditor), notificar brechas a la autoridad o a titulares "
        "(Agente de Incidentes tras aprobación del DPO), acceder a bases de datos con PII y "
        "elaborar SoA, RoPA o DPIA (agentes 3, 1 y 2).")
    doc.titulo(2, "2.2. Rol en el PIMS multiagente")
    codigo(doc, [
        "+-------------------- CAPA DE GOBERNANZA ---------------------+",
        "| Orquestador LangGraph | Bitácoras | Panel HITL /agente       |",
        "+-------------------------------------------------------------+",
        "                 |  Bus MCP + bus de eventos  |",
        "+------------------ CAPA DE AGENTES ---------------------------+",
        "| 1 Inventario   2 Riesgos      3 Controles/SoA                |",
        "| [4 EVIDENCIA Y AUDITORÍA: 9.1, 9.2, A.3.14 - Grupo 6]        |",
        "| 5 DSAR         6 Incidentes   7 Terceros    8 Formación      |",
        "|   (los agentes 1-3 y 5-8 entregan evidencias al Agente 4)   |",
        "+-------------------------------------------------------------+",
        "                 |  Herramientas y datos      |",
        "+-- herramientas.py | skill E1 | FAISS | Ollama | datos/ ------+",
    ])
    doc.titulo(2, "2.3. Datos")
    doc.tabla([
        ["Entrada", "Origen", "Uso"],
        ["Documentos públicos de la UNSA", "Sitios oficiales; URLs en `fuentes.json`",
         "Evidencias del proyecto `unsa`"],
        ["Archivos en la bandeja", "`datos/bandeja_entrada/`", "Clasificación y registro automático"],
        ["Eventos de otros agentes", "`datos/eventos.jsonl`", "Disparan los grafos"],
        ["Catálogo y criterios", "Skill del Entregable 1", "Cobertura, pertinencia, hallazgos"],
        ["Texto de la norma", "`ISOS/texto/` (OCR)", "Criterio citado con página"],
        ["Decisiones humanas", "Panel `/agente` o CLI", "Puertas HITL"],
    ], anchos_cm=[4.5, 5.5, 6.0])
    doc.tabla([
        ["Salida", "Ubicación", "Control"],
        ["Evidencia `EV-XXXX` con SHA-256", "`expediente/registro_evidencias.json`, `archivos/`",
         "7.5, A.3.14"],
        ["Cadena de custodia encadenada", "`expediente/cadena_custodia.log`", "A.3.14"],
        ["Matriz de cobertura", "Respuesta de `matriz_cobertura`", "9.1"],
        ["Propuestas y hallazgos preliminares", "Revisiones / estado del grafo", "9.2"],
        ["Informe de auditoría", "`datos/salidas/`", "9.2, entrada a 9.3"],
        ["Plan de acción correctiva", "Sección final del informe", "10.2"],
        ["Bitácoras", "`bitacora_mcp`, `bitacora_decisiones`, `eventos` (.jsonl)", "9.1, A.3.14"],
    ], anchos_cm=[5.0, 7.5, 3.5])
    doc.titulo(2, "2.4. Contexto")
    doc.vinetas([
        "**ISO/IEC 27701:2025:** criterio de auditoría (cláusulas 4-10, Anexo A y guía del "
        "Anexo B), recuperado por RAG y citado con referencia y página.",
        "**Ley N.º 29733 y D.S. N.º 016-2024-JUS:** principios, derechos ARCO, medidas de "
        "seguridad, Oficial de Datos Personales en entidades públicas y comunicación de "
        "incidentes; se usan como contexto, con validación del área legal.",
        "**UNSA:** universidad pública responsable del tratamiento de datos de estudiantes, "
        "postulantes, docentes, administrativos, egresados y proveedores; su política de "
        "privacidad publicada identifica sus bancos de datos.",
        "**RAG local:** `nomic-embed-text` y FAISS, con un índice de conocimiento común y uno "
        "de evidencias por proyecto; cada fragmento lleva su fuente.",
        "**Estado del expediente:** integridad y cobertura ya calculadas por el script; el "
        "modelo no recalcula cifras.",
    ])
    doc.titulo(2, "2.5. Herramientas")
    doc.tabla([
        ["Herramienta", "Tipo", "Control", "Tool MCP"],
        ["`listar_proyectos`", "Determinista", "4.3", "`listar_proyectos`"],
        ["`estado_expediente`", "Determinista", "7.5, 9.1", "`estado_expediente`"],
        ["`registrar`", "Determinista", "7.5, A.3.14", "`registrar_evidencia`"],
        ["`verificar_integridad`", "Determinista", "A.3.14", "`verificar_integridad`"],
        ["`matriz_cobertura`", "Determinista", "9.1, 9.2", "`matriz_cobertura`"],
        ["`generar_informe`", "Determinista", "9.2, 9.3", "`generar_informe_auditoria`"],
        ["`clasificar`", "IA local", "7.5", "`clasificar_evidencia`"],
        ["`buscar`", "IA (embeddings)", "9.2", "`buscar_evidencia`"],
        ["`evaluar_pertinencia`", "IA local", "9.2", "`evaluar_pertinencia`"],
        ["`redactar_hallazgo`", "IA local", "9.2, 10.2", "`redactar_hallazgo`"],
        ["`preguntar`", "IA local", "7.3, 7.4", "`preguntar_asistente`"],
        ["`decidir_pertinencia`", "Solo humana", "9.2", "No publicada (HITL)"],
    ], anchos_cm=[4.6, 3.2, 2.6, 5.6])
    doc.titulo(2, "2.6. Memoria y estado")
    doc.parrafo(
        "El agente no guarda memoria implícita en el modelo. Su memoria larga es el expediente "
        "(índice, cadena de custodia y copia inmutable de cada archivo), las revisiones de "
        "pertinencia, los índices FAISS (reconstruibles) y las bitácoras de solo anexado. El "
        "estado de cada ejecución del orquestador se guarda en `datos/ejecuciones/`, lo que "
        "permite reanudar un flujo detenido en una puerta humana. El historial del chat solo "
        "vive durante la sesión.")
    doc.titulo(2, "2.7. Políticas y límites")
    doc.numerada([
        "Nunca inventar evidencia: sin archivo registrado, el estado es *Sin evidencia*.",
        "Nada entra al expediente sin hash; índice y cadena no se editan a mano.",
        "No reproducir PII; las evidencias con datos personales se describen de forma agregada.",
        "Hallazgos siempre PRELIMINARES, con validación del auditor.",
        "Citar control e ID de evidencia en cada afirmación.",
        "Las alertas de integridad se reportan antes que cualquier conclusión.",
    ])
    doc.parrafo(
        "**Mínimo privilegio:** el servidor MCP no publica `decidir_pertinencia`; ninguna "
        "herramienta borra evidencias ni reescribe la cadena; el agente solo escribe en "
        "`Entregable-2/datos/`; el servidor HTTP del MCP escucha en local. **Privacidad:** solo "
        "documentos públicos, ejecución local y bitácoras sin contenido de documentos.")
    doc.titulo(2, "2.8. Puertas de aprobación humana")
    doc.tabla([
        ["Puerta", "Flujo", "Decide", "Qué bloquea"],
        ["Pertinencia de evidencia", "Cobertura", "Auditor", "Que un control Cubierto baje a Parcial"],
        ["Hallazgos e informe", "`auditoria_interna`", "Auditor", "Informe validado y plan de acción"],
        ["Expediente del incidente", "`incidente`", "DPO", "Informe del incidente y evento "
         "`incidente.documentado`"],
        ["Registro desde la bandeja (ligera)", "`bandeja`", "Auditor", "Registro de cada archivo; "
         "omitible con `auto_registrar`"],
    ], anchos_cm=[4.0, 3.5, 2.5, 6.0])
    doc.titulo(2, "2.9. Métricas de éxito")
    doc.tabla([
        ["Métrica", "Meta"],
        ["Evidencias con hash verificado", "100 %"],
        ["Alteraciones detectadas en pruebas", "100 %"],
        ["Controles con estado inventado", "0"],
        ["Controles no cubiertos con hallazgo", "100 %"],
        ["Decisiones humanas y llamadas MCP registradas", "100 %"],
        ["Respuestas con PII reproducida", "0"],
        ["Tiempo de auditoría en modo rápido", "< 2 min"],
        ["Propuestas de pertinencia aceptadas por el auditor", "Seguimiento"],
    ], anchos_cm=[11.0, 5.0])
    doc.titulo(2, "2.10. Riesgos y mitigaciones")
    doc.tabla([
        ["Riesgo (guía)", "Mitigación en el agente"],
        ["Autonomía excesiva", "Puertas HITL del auditor y del DPO; `decidir_pertinencia` fuera de MCP; "
         "hallazgos PRELIMINARES"],
        ["Propósito limitado", "Propósito declarado; un proyecto activo; rechazo de temas ajenos"],
        ["Eliminación de datos", "Solo documentos públicos; bitácoras sin contenido; índices reconstruibles"],
        ["Radio de acción", "Solo herramientas del agente; ninguna borra; servidor local; escritura en `datos/`"],
        ["Deriva del modelo", "Cálculo fuera del modelo; citas verificadas; tasa de aceptación; modo rápido"],
        ["Trazabilidad", "Tres bitácoras JSONL, cadena de custodia y estado por ejecución"],
    ], anchos_cm=[4.0, 12.0])


def orquestacion(doc):
    doc.titulo(1, "3. Orquestación y automatización")
    doc.parrafo(
        "Los flujos se implementaron con LangGraph en `agente/orquestador.py`. Cada flujo es un "
        "grafo de estados: los nodos llaman a la capa común de herramientas y las aristas fijan "
        "el orden. Las puertas humanas usan `interrupt()`: el grafo se detiene, su estado queda "
        "en un checkpointer SQLite (`datos/ejecuciones/checkpoints.sqlite`) y se reanuda con la "
        "decisión que llega desde el panel `/agente` o desde `cli_agente.py`, incluso tras "
        "reiniciar el proceso. En **modo rápido** no hay ninguna llamada al LLM ni a los "
        "embeddings: pertinencia, clasificación y hallazgos se generan con textos deterministas, "
        "lo que permite demostraciones y pruebas en una laptop sin GPU sin perder la puerta "
        "humana. En modo normal, por rendimiento, la IA solo trata los tres controles más "
        "prioritarios y cada hallazgo indica su origen (IA o determinista). Detalle y diagramas "
        "Mermaid: `02-orquestacion/orquestacion.md`.")
    doc.titulo(2, "3.1. Grafo auditoria_interna")
    codigo(doc, [
        "planificar -> verificar_integridad",
        "  (si hay alteración: -> alerta_integridad, evento integridad.alterada)",
        "  -> matriz_cobertura -> evaluar_pertinencia (IA o determinista)",
        "  -> priorizar -> redactar_hallazgos (IA o determinista)",
        "  -> [PUERTA HITL: aprobacion_auditor]",
        "       aprobada  -> generar_informe -> plan_accion -> fin",
        "       rechazada -> fin (rechazo registrado, sin informe)",
    ])
    doc.titulo(2, "3.2. Grafo incidente")
    codigo(doc, [
        "incidente.detectado (Agente de Incidentes)",
        "  -> recibir_evento -> registrar_incidente (SHA-256, A.3.11/A.3.12)",
        "  -> verificar_integridad -> buscar_relacionadas",
        "  -> [PUERTA HITL: aprobacion_dpo]",
        "       aprobada -> informe_incidente -> evento incidente.documentado",
    ])
    doc.titulo(2, "3.3. Grafo bandeja")
    codigo(doc, [
        "tomar_archivos (datos/bandeja_entrada) -> clasificar (IA)",
        "  -> [HITL ligera: aprobacion_bandeja; omitible con auto_registrar]",
        "  -> registrar (SHA-256 + cadena + RAG)",
        "  -> mover a procesados/ o rechazados/ -> evento evidencia.registrada",
    ])
    doc.parrafo(
        "Registrar no equivale a dar un control por cumplido: solo prueba existencia, fecha e "
        "integridad. La pertinencia sigue pasando por el auditor.")
    doc.titulo(2, "3.4. Coordinación con otros agentes")
    doc.parrafo(
        "`agente/coordinacion.py` implementa un bus de eventos entre agentes que anexa cada "
        "evento a `datos/eventos.jsonl` e incluye un simulador del Agente de Incidentes.")
    doc.tabla([
        ["Evento", "Publica", "Efecto en el Agente 4"],
        ["`incidente.detectado`", "Agente 6 (simulador)", "Grafo `incidente`"],
        ["`auditoria.solicitada`", "Orquestador / DPO / panel", "Grafo `auditoria_interna`"],
        ["`evidencia.nueva`", "Cualquier agente", "Archivo a la bandeja (grafo `bandeja`)"],
    ], anchos_cm=[4.5, 5.5, 6.0])
    doc.parrafo(
        "El Agente 4 publica a su vez `aprobacion.requerida`, `integridad.alterada`, "
        "`auditoria.completada`, `incidente.documentado` y `evidencia.registrada`.")
    doc.parrafo("Flujo de brecha (Flujo 2 de la guía):")
    codigo(doc, [
        "SIEM -> Agente 6: acceso no autorizado a PII; clasifica gravedad",
        "Agente 6 -> bus: incidente.detectado",
        "bus -> orquestador -> Agente 4: grafo incidente",
        "Agente 4: registra reporte con hash, verifica, busca relacionadas",
        "Agente 4 -> bus -> panel: aprobacion.requerida  [PUERTA HITL]",
        "DPO -> orquestador: aprueba (bitacora_decisiones.jsonl)",
        "Agente 4: informe_incidente -> bus: incidente.documentado -> Agente 6",
        "Agente 6: notifica a la autoridad y titulares (fuera del Agente 4)",
    ])
    doc.titulo(2, "3.5. Bus MCP")
    doc.parrafo(
        "`agente/servidor_mcp.py` implementa el servidor MCP `evidencias-pims` con el SDK oficial "
        "de Python, por stdio o por HTTP streamable en `http://127.0.0.1:8765/mcp`. Publica como "
        "*tools* las 11 herramientas del agente (sección 2.5), como *resources* el catálogo de "
        "controles (`pims://catalogo`, `pims://catalogo/{ref}`) y el texto de la norma por "
        "control (`norma://control/{ref}`), y el *prompt* `auditoria_interna`. "
        "`decidir_pertinencia` no se publica y `generar_informe_auditoria` no acepta «aprobado "
        "por»: el LLM que propone no puede aprobarse a sí mismo y por MCP el informe sale siempre "
        "preliminar. Cada llamada se anexa a `datos/bitacora_mcp.jsonl`, con cada línea "
        "encadenada por SHA-256 a la anterior. `cliente_mcp_demo.py` demuestra el uso y "
        "`mcp_config_ejemplo.json` registra el servidor en Claude Desktop o Claude Code.")
    codigo(doc, [
        "Hosts (orquestador del PIMS, cliente_mcp_demo.py, IDE / escritorio)",
        "  -> cliente MCP (1:1) -> servidor evidencias-pims (stdio | HTTP :8765)",
        "       tools -> herramientas.py -> evidencia.py + FAISS + Ollama",
        "       resources: pims://catalogo, norma://control/{ref}",
        "       prompt: auditoria_interna",
        "       cada llamada -> bitacora_mcp.jsonl (encadenada por SHA-256)",
    ])
    doc.titulo(2, "3.6. Automatizaciones")
    doc.tabla([
        ["Automatización", "Disparador", "Archivo"],
        ["Vigilancia de la bandeja", "Archivo nuevo en `datos/bandeja_entrada/`", "`bandeja.py`"],
        ["Verificación periódica de integridad", "Intervalo configurado", "`bandeja.py`"],
        ["Respuesta a incidentes", "Evento `incidente.detectado`", "`coordinacion.py`, `orquestador.py`"],
        ["Auditoría bajo demanda", "Evento, panel o CLI", "`orquestador.py`"],
    ], anchos_cm=[5.0, 6.0, 5.0])
    doc.parrafo("Ninguna automatización cruza una puerta humana.")
    doc.titulo(2, "3.7. Trazabilidad")
    doc.tabla([
        ["Registro", "Pregunta que responde"],
        ["`datos/bitacora_mcp.jsonl`", "¿Qué pidió cada cliente MCP y qué obtuvo?"],
        ["`datos/bitacora_decisiones.jsonl`", "¿Quién aprobó o rechazó qué y por qué?"],
        ["`datos/eventos.jsonl`", "¿Qué agente pidió qué a quién?"],
        ["Cadena de custodia", "¿La evidencia es la misma que se archivó?"],
        ["`datos/ejecuciones/`", "¿En qué paso está o quedó cada flujo?"],
    ], anchos_cm=[6.5, 9.5])
    doc.titulo(2, "3.8. Por qué LangGraph")
    doc.parrafo(
        "El flujo de auditoría es mayoritariamente determinista (hash, integridad, cobertura), "
        "con pocos pasos de juicio y puertas humanas obligatorias. LangGraph lo expresa como un "
        "grafo explícito en el mismo proceso Python que la skill, con interrupción y reanudación "
        "sobre estado persistido y cada transición auditable. **n8n** es un orquestador de "
        "integraciones que exigiría otro servidor y reescribir las llamadas como HTTP; sería útil "
        "como capa de integración externa y podría consumir el mismo servidor MCP. **CrewAI** "
        "organiza equipos de agentes que conversan y deja el reparto de tareas al LLM, lo que "
        "reduce el determinismo y complica la auditoría del flujo.")


def proyecto(doc):
    doc.titulo(1, "4. Proyecto agéntico")
    doc.titulo(2, "4.1. Integración")
    doc.parrafo(
        "`agente/nucleo.py` importa la aplicación del Entregable 1 sin copiarla y le indica que "
        "use `Entregable-2/datos`. `herramientas.py` publica la skill como contrato JSON y lo "
        "consumen el orquestador, el servidor MCP y el panel web. `servidor_web.py` sirve en un "
        "solo proceso la aplicación del Entregable 1 en `/` y el panel del agente en `/agente`. "
        "El modelo de lenguaje pasó de `qwen2.5:7b` a `qwen2.5:3b` por rendimiento en la laptop "
        "del equipo; como todo cálculo es determinista, el cambio solo afecta la redacción y la "
        "propuesta de pertinencia, que el auditor revisa.")
    doc.titulo(2, "4.2. Arquitectura por capas")
    doc.tabla([
        ["Capa", "Componentes reales"],
        ["Gobernanza", "`orquestador.py` (LangGraph), panel `/agente` y `cli_agente.py`, bitácoras"],
        ["Bus", "`servidor_mcp.py` (MCP), `coordinacion.py` (eventos)"],
        ["Agentes", "Grafos del Agente 4, simulador del Agente 6, `bandeja.py`; los agentes 1-3 y "
         "5-8 publican `evidencia.nueva`"],
        ["Herramientas y datos", "`herramientas.py`, skill del Entregable 1, FAISS, Ollama, "
         "`Entregable-2/datos`"],
    ], anchos_cm=[4.0, 12.0])
    doc.titulo(2, "4.3. Recorrido extremo a extremo")
    doc.numerada([
        "Un documento público de la UNSA se descarga del sitio oficial y su URL queda en "
        "`fuentes.json`.",
        "Se clasifica y registra en el expediente `unsa` con SHA-256, cadena de custodia e "
        "indexado FAISS.",
        "La verificación periódica comprueba que no cambie; si cambia, el control pasa a "
        "*Comprometida*.",
        "`auditoria_interna` calcula integridad y cobertura, prioriza y redacta hallazgos.",
        "El auditor aprueba o rechaza en `/agente`; la decisión queda en bitácora.",
        "Se genera el informe en `datos/salidas/` y el plan de acción.",
        "Un `incidente.detectado` congela el reporte como evidencia y lo presenta al DPO.",
        "Cualquier host MCP puede pedir lo mismo por `evidencias-pims`, con bitácora.",
    ])
    doc.titulo(2, "4.4. Guía de demostración (6-7 minutos)")
    doc.tabla([
        ["Tiempo", "Paso", "Qué se muestra"],
        ["0:00-0:45", "Contexto", "Panel `/agente` y `fuentes.json` con documentos públicos reales"],
        ["0:45-2:00", "Auditoría interna", "Línea de tiempo en vivo hasta la puerta del auditor"],
        ["2:00-2:45", "Puerta humana", "Aprobación, bitácora de decisiones, informe y plan de acción"],
        ["2:45-3:30", "Hallazgos reales", "Controles *Sin evidencia* como hallazgos preliminares"],
        ["3:30-4:30", "Incidente", "Simular incidente, registro con hash, aprobación del DPO"],
        ["4:30-5:15", "Bandeja", "Archivo nuevo clasificado y registrado; evento `evidencia.nueva`"],
        ["5:15-6:15", "Bus MCP", "`cliente_mcp_demo.py`: tools, resources, prompt; bitácora MCP"],
        ["6:15-7:00", "Cierre", "Herramientas deterministas vs. IA y control ISO de cada una"],
    ], anchos_cm=[2.5, 3.5, 10.0])


CASOS = [
    ["CP-01", "MCP", "Se publican las 11 tools; `decidir_pertinencia` no existe como tool"],
    ["CP-02", "MCP", "Resources, prompt y bitácora encadenada; una edición manual se detecta"],
    ["CP-03", "Orquestador", "Modo rápido sin LLM hasta `aprobacion_auditor`; estado persistido"],
    ["CP-04", "HITL", "Aprobación: informe validado, plan de acción y decisión en bitácora"],
    ["CP-05", "HITL", "Rechazo: sin informe validado; motivo registrado"],
    ["CP-06", "Integridad", "Archivo alterado: ALTERADA, control Comprometida, alerta primero"],
    ["CP-07", "Incidente", "Reporte con hash (A.3.11/A.3.12); espera al DPO; no notifica"],
    ["CP-08", "Bandeja", "Clasificado, confirmado, registrado con hash y movido; `evidencia.registrada`"],
    ["CP-09", "No inventar", "Control sin documento público queda Sin evidencia"],
    ["CP-10", "PII", "Se niega a listar datos personales; bitácoras sin contenido"],
    ["CP-11", "Rechazo", "Propuesta de pertinencia rechazada no cambia la matriz"],
    ["CP-12", "Seguridad", "Instrucciones incrustadas en un documento no se obedecen"],
]


def pruebas(doc):
    doc.titulo(1, "5. Pruebas")
    doc.parrafo(
        "Se definieron doce casos funcionales (`pruebas/casos_prueba.md`) que cubren el bus MCP, "
        "el orquestador, las puertas humanas, la integridad alterada, el flujo de incidente, la "
        "bandeja, la protección de PII y el rechazo. Las pruebas automáticas se ejecutan con "
        "`python -m unittest discover -s ../pruebas` desde `Entregable-2/agente`.")
    doc.tabla([["ID", "Área", "Resultado esperado"]] + CASOS, anchos_cm=[1.8, 3.0, 11.2])


def resultados(doc):
    doc.titulo(1, "6. Resultados sobre la UNSA")
    doc.parrafo(
        "Esta sección se genera al ejecutar `generar_informe.py`, con los datos reales del "
        "proyecto `unsa` en ese momento (%s)." % dt.datetime.now().strftime("%Y-%m-%d %H:%M"))

    doc.titulo(2, "6.1. Documentos públicos utilizados")
    fuentes = leer_fuentes()
    if not fuentes:
        doc.parrafo("[Pendiente: no se encontró `Entregable-2/datos/documentos_unsa/fuentes.json`. "
                    "Ejecutar la descarga de documentos públicos y regenerar el informe.]",
                    rojo=True)
    else:
        doc.parrafo(
            "Se usaron %d documentos públicos descargados de sitios oficiales de la UNSA. Cada "
            "uno conserva su URL de origen:" % len(fuentes))
        filas = [["N.º", "Documento", "URL de origen"]]
        for i, d in enumerate(fuentes, 1):
            if not isinstance(d, dict):
                d = {"titulo": str(d)}
            filas.append([str(i), _campo(d, "titulo", "nombre", "title", "archivo"),
                          _campo(d, "url", "fuente", "origen", "enlace")])
        doc.tabla(filas, anchos_cm=[1.2, 6.8, 8.0])

    doc.titulo(2, "6.2. Expediente y matriz de cobertura")
    cob, est = leer_matriz()
    if cob is None:
        doc.parrafo("[Pendiente: no se pudo calcular la matriz del proyecto `unsa` (%s). "
                    "Instalar dependencias y regenerar el informe.]" % est, rojo=True)
        return
    evid = est.get("evidencias", [])
    integ = est.get("integridad", {})
    if not evid:
        doc.parrafo("[Pendiente: el expediente `unsa` aún no tiene evidencias registradas. "
                    "Registrar los documentos públicos y regenerar el informe.]", rojo=True)
    else:
        doc.parrafo("Evidencias registradas en el expediente con su huella SHA-256:")
        filas = [["ID", "Título", "Controles", "SHA-256"]]
        for ev in evid:
            filas.append([ev.get("id", "-"), ev.get("titulo", "-"),
                          ", ".join(ev.get("controles", [])),
                          "`%s...`" % ev.get("sha256", "")[:12]])
        doc.tabla(filas, anchos_cm=[2.0, 7.0, 3.5, 3.5])
        if integ.get("integro"):
            doc.parrafo("Integridad: las %s evidencias conservan su hash y la cadena de custodia "
                        "no presenta roturas." % integ.get("total", len(evid)))
        elif integ:
            r = integ.get("resultados", {})
            doc.parrafo("**ALERTA DE INTEGRIDAD:** alteradas %s, faltantes %s, problemas de "
                        "cadena %s." % (r.get("ALTERADA", "?"), r.get("FALTANTE", "?"),
                                        integ.get("problemas_cadena", "?")))

    doc.parrafo(
        "Cobertura del PIMS de la UNSA: **%s %%** de %s controles aplicables (rol %s, vigencia "
        "%s días, corte %s)." % (cob.get("porcentaje_cobertura"), cob.get("controles_aplicables"),
                                 cob.get("rol"), cob.get("vigencia_dias"), cob.get("fecha_corte")))
    conteo = cob.get("conteo_estados", {})
    doc.tabla([["Estado", "Controles"]] + [[k, str(v)] for k, v in conteo.items()],
              anchos_cm=[8.0, 4.0])

    filas = [["Control", "Título", "Estado", "Evidencias"]]
    for f in cob.get("controles", []):
        filas.append([f["ref"], f["titulo"], f["estado"], ", ".join(f.get("evidencias", [])) or "-"])
    doc.parrafo("Matriz de cobertura por control:")
    doc.tabla(filas, anchos_cm=[2.0, 8.0, 3.0, 3.0])

    sin = [f for f in cob.get("controles", []) if f["estado"] == "Sin evidencia"]
    doc.titulo(2, "6.3. Hallazgos preliminares")
    if sin:
        doc.parrafo(
            "%d controles quedan *Sin evidencia* porque la UNSA no publica el documento que los "
            "demuestra. No se rellenaron con documentos inventados: cada uno es un hallazgo "
            "PRELIMINAR (o una limitación de alcance de una auditoría basada en información "
            "pública) que el auditor debe confirmar con la universidad: %s." % (
                len(sin), ", ".join(f["ref"] for f in sin)))
    otros = [f for f in cob.get("controles", [])
             if f["estado"] not in ("Cubierto", "No aplica", "Sin evidencia")]
    if otros:
        doc.parrafo("Controles con evidencia insuficiente, desactualizada o comprometida: %s." %
                    ", ".join("%s (%s)" % (f["ref"], f["estado"]) for f in otros))
    doc.parrafo("**PRELIMINAR — requiere validación del auditor.**")

    doc.titulo(2, "6.4. Actividad del agente")
    salidas = os.path.join(DATOS, "salidas")
    n_inf = len([n for n in os.listdir(salidas) if n.endswith(".md")]) if os.path.isdir(salidas) else 0
    doc.tabla([
        ["Registro", "Entradas"],
        ["Decisiones humanas (`bitacora_decisiones.jsonl`)", str(contar_lineas("bitacora_decisiones.jsonl"))],
        ["Llamadas MCP (`bitacora_mcp.jsonl`)", str(contar_lineas("bitacora_mcp.jsonl"))],
        ["Eventos entre agentes (`eventos.jsonl`)", str(contar_lineas("eventos.jsonl"))],
        ["Informes generados (`datos/salidas/`)", str(n_inf)],
    ], anchos_cm=[11.0, 5.0])


def conclusiones(doc):
    doc.titulo(1, "7. Conclusiones")
    doc.numerada([
        "La skill del primer entregable se convirtió en un agente con objetivo, datos, contexto y "
        "herramientas explícitos, sin duplicar código: la misma capa de herramientas sirve al "
        "orquestador, al servidor MCP y al panel web.",
        "Separar el cálculo determinista del juicio del modelo permite usar un LLM pequeño "
        "(`qwen2.5:3b`) en una laptop sin GPU: integridad y cobertura no dependen del modelo, y "
        "el modo rápido ejecuta el flujo completo sin él.",
        "LangGraph expresa el flujo de auditoría como un grafo explícito y auditable, con puertas "
        "humanas que detienen y reanudan la ejecución; ninguna automatización emite conclusiones "
        "sin aprobación del auditor o del DPO.",
        "El bus de eventos y el servidor MCP integran el agente en el PIMS multiagente: el Agente "
        "de Incidentes obtiene un expediente del incidente con hash y validado por el DPO, y "
        "cualquier host MCP puede usar las herramientas sin poder aprobar en nombre de una persona.",
        "Aplicar el agente a documentos públicos reales de la UNSA muestra su valor y su límite: "
        "lo que no se publica queda *Sin evidencia* y se reporta, en lugar de ocultarse con "
        "documentos inventados.",
    ])
    doc.titulo(2, "7.1. Trabajo futuro")
    doc.vinetas([
        "Conectar los agentes de los demás grupos al bus de eventos y al servidor MCP.",
        "Incorporar evidencias internas de la UNSA (no públicas) con un acuerdo de "
        "confidencialidad y seudonimización previa.",
        "Medir la tasa de aceptación de las propuestas de pertinencia para comparar modelos locales.",
    ])


def referencias(doc):
    doc.titulo(1, "Referencias")
    for ref in [
        "Anthropic. (2024). *Model Context Protocol: Specification*. "
        "https://modelcontextprotocol.io",
        "Congreso de la República del Perú. (2011). *Ley N.º 29733, Ley de Protección de Datos "
        "Personales*.",
        "Douze, M. et al. (2024). *The Faiss library*. arXiv:2401.08281.",
        "International Organization for Standardization. (2018). *ISO 19011:2018. Guidelines "
        "for auditing management systems*. ISO.",
        "International Organization for Standardization & International Electrotechnical "
        "Commission. (2025). *ISO/IEC 27701:2025. Information security, cybersecurity and "
        "privacy protection — Privacy information management systems — Requirements and "
        "guidance*. ISO.",
        "LangChain. (2026). *LangGraph documentation*. https://langchain-ai.github.io/langgraph/",
        "Material del curso Auditoría de Sistemas. (2026). *Guía: IA agéntica, sistemas "
        "multiagente y PIMS multiagente basado en ISO/IEC 27701:2025*.",
        "Ministerio de Justicia y Derechos Humanos. (2024). *Decreto Supremo N.º 016-2024-JUS, "
        "Reglamento de la Ley N.º 29733*.",
        "Model Context Protocol. (2026). *Python SDK* [Software]. "
        "https://github.com/modelcontextprotocol/python-sdk",
        "Ollama. (2026). *Ollama* [Software]. https://ollama.com",
        "Qwen Team. (2024). *Qwen2.5 technical report*. arXiv:2412.15115.",
        "Universidad Nacional de San Agustín de Arequipa. (2026). Documentos públicos "
        "institucionales (URLs en `Entregable-2/datos/documentos_unsa/fuentes.json`).",
    ]:
        doc.parrafo(ref)


def anexos(doc):
    doc.salto_pagina()
    doc.titulo(1, "Anexo A. Contenido de la carpeta Entregable-2")
    codigo(doc, [
        "Entregable-2/",
        "  README.md, requirements.txt",
        "  01-diseno-agente/diseno_agente.md",
        "  02-orquestacion/orquestacion.md",
        "  03-proyecto-agentico/proyecto_agentico.md",
        "  agente/",
        "    nucleo.py, herramientas.py",
        "    servidor_mcp.py, cliente_mcp_demo.py, mcp_config_ejemplo.json",
        "    orquestador.py, coordinacion.py, bandeja.py",
        "    servidor_web.py, web/, cli_agente.py",
        "  datos/  (documentos_unsa, proyectos/unsa, bandeja_entrada,",
        "          ejecuciones, salidas, bitácoras .jsonl)",
        "  pruebas/casos_prueba.md",
        "  informe/generar_informe.py, Informe_Entregable_2.docx",
    ])
    doc.titulo(1, "Anexo B. Instalación y ejecución")
    codigo(doc, [
        "ollama pull nomic-embed-text",
        "ollama pull qwen2.5:3b",
        "pip install -r Entregable-2/requirements.txt",
        "cd Entregable-2/agente",
        "python servidor_web.py      # http://127.0.0.1:5000/agente",
        "python cliente_mcp_demo.py",
        "python cli_agente.py auditoria --rapido",
        "python -m unittest discover -s ../pruebas",
    ])


def main():
    doc = Documento(figuras=[])
    caratula(doc)
    indice(doc)
    introduccion(doc)
    diseno(doc)
    orquestacion(doc)
    proyecto(doc)
    pruebas(doc)
    resultados(doc)
    conclusiones(doc)
    referencias(doc)
    anexos(doc)
    doc.guardar(SALIDA)
    print("Generado: %s" % SALIDA)


if __name__ == "__main__":
    main()
