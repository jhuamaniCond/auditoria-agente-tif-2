# -*- coding: utf-8 -*-
"""
docx_builder.py
===============================================================================
Motor comun para generar los documentos Word del laboratorio con python-docx.
Lo usan generar_word_informe.py y generar_word_plantilla.py.

Requisito:
    pip install python-docx

Estilo academico aplicado a todo el documento:
  - Times New Roman, 12 pt en ABSOLUTAMENTE todo el texto: cuerpo, titulos de
    cualquier nivel, listas, tablas y leyendas de figura.
  - Sin colores. El unico color permitido es el ROJO, reservado para las
    indicaciones internas / texto pendiente que se elimina del informe final
    (parametro rojo=True).
  - Tablas sin formato: solo lineas simples negras de 0.5 pt, sin sombreados,
    sin bandas de color y sin estilos decorativos de Word.
  - Pagina A4, margenes de 2.5 cm, interlineado sencillo, texto justificado.

Marcado en linea admitido dentro de los textos:
    **negrita**      -> negrita
    *cursiva*        -> cursiva
    `monoespaciado`  -> Courier New (equivalente a \\texttt de LaTeX)
===============================================================================
"""

import os
import re

from docx import Document as _DocumentoWord
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# --------------------------------------------------------------------------
# Parametros de estilo
# --------------------------------------------------------------------------
FUENTE = "Times New Roman"
FUENTE_MONO = "Courier New"
TAMANO_PT = 12                        # tamano unico de letra en todo el documento
ROJO = RGBColor(0xC0, 0x00, 0x00)     # unico color permitido, solo indicaciones
NEGRO = RGBColor(0x00, 0x00, 0x00)

MARGEN_CM = 2.5
ANCHO_UTIL_CM = 21.0 - 2 * MARGEN_CM  # 16 cm de ancho aprovechable en A4

_PATRON_MARCADO = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)", re.DOTALL)


def _trocear(texto):
    """Divide un texto con marcado en linea en tramos
    (contenido, negrita, cursiva, mono)."""
    tramos = []
    for parte in _PATRON_MARCADO.split(texto):
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**") and len(parte) > 4:
            tramos.append((parte[2:-2], True, False, False))
        elif parte.startswith("*") and parte.endswith("*") and len(parte) > 2:
            tramos.append((parte[1:-1], False, True, False))
        elif parte.startswith("`") and parte.endswith("`") and len(parte) > 2:
            tramos.append((parte[1:-1], False, False, True))
        else:
            tramos.append((parte, False, False, False))
    return tramos


def _escribir(parrafo, texto, negrita=False, cursiva=False, rojo=False):
    """Vuelca un texto con marcado dentro de un parrafo ya creado."""
    for contenido, b, i, mono in _trocear(texto):
        run = parrafo.add_run(contenido)
        run.bold = negrita or b
        run.italic = cursiva or i
        run.font.name = FUENTE_MONO if mono else FUENTE
        run.font.size = Pt(TAMANO_PT)
        run.font.color.rgb = ROJO if rojo else NEGRO
        # La fuente tambien debe fijarse para los caracteres de Asia oriental,
        # de lo contrario Word puede sustituirla en algunos equipos.
        run._element.rPr.rFonts.set(qn("w:eastAsia"),
                                    FUENTE_MONO if mono else FUENTE)
    return parrafo


# --------------------------------------------------------------------------
# Documento
# --------------------------------------------------------------------------
class Documento(object):
    """Envoltorio de python-docx con el estilo academico ya configurado."""

    def __init__(self, figuras=None):
        self.doc = _DocumentoWord()
        self._configurar_pagina()
        self._configurar_estilos()
        # Registro de figuras: permite citar "Figura N" antes de insertarla,
        # igual que \ref{} en LaTeX. Se recibe la lista ordenada de etiquetas.
        self._figuras = {}
        if figuras:
            for indice, etiqueta in enumerate(figuras, start=1):
                self._figuras[etiqueta] = indice

    # ------------------------------------------------------------ preparacion
    def _configurar_pagina(self):
        """A4 vertical con margenes de 2.5 cm en los cuatro lados."""
        seccion = self.doc.sections[0]
        seccion.orientation = WD_ORIENT.PORTRAIT
        seccion.page_width = Cm(21.0)
        seccion.page_height = Cm(29.7)
        seccion.top_margin = Cm(MARGEN_CM)
        seccion.bottom_margin = Cm(MARGEN_CM)
        seccion.left_margin = Cm(MARGEN_CM)
        seccion.right_margin = Cm(MARGEN_CM)

    def _configurar_estilos(self):
        """Fija Times New Roman 12 pt como base y homogeneiza los titulos.

        Los tres niveles de titulo comparten el mismo tamano de 12 pt y se
        distinguen unicamente por la negrita, tal como exige la plantilla
        academica (equivale al \\titleformat{...}{\\fontsize{12}{14}\\bfseries}
        del archivo LaTeX original)."""
        normal = self.doc.styles["Normal"]
        normal.font.size = Pt(TAMANO_PT)
        normal.font.color.rgb = NEGRO
        self._fijar_fuente(normal)

        parrafo = normal.paragraph_format
        parrafo.space_before = Pt(0)
        parrafo.space_after = Pt(6)
        parrafo.line_spacing = 1.0            # interlineado sencillo
        parrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        for nivel, antes in ((1, 12), (2, 10), (3, 8)):
            estilo = self.doc.styles["Heading %d" % nivel]
            estilo.base_style = normal
            estilo.font.size = Pt(TAMANO_PT)   # sin escalado por nivel
            estilo.font.bold = True
            estilo.font.italic = False
            estilo.font.color.rgb = NEGRO      # Word los pone azules por defecto
            self._fijar_fuente(estilo)
            formato = estilo.paragraph_format
            formato.alignment = WD_ALIGN_PARAGRAPH.LEFT
            formato.space_before = Pt(antes)
            formato.space_after = Pt(6)
            formato.line_spacing = 1.0
            formato.keep_with_next = True

        # Estilo de las leyendas de figura: mismo cuerpo de 12 pt, centrado.
        leyenda = self.doc.styles["Caption"]
        leyenda.base_style = normal
        leyenda.font.size = Pt(TAMANO_PT)
        leyenda.font.italic = False
        leyenda.font.bold = False
        leyenda.font.color.rgb = NEGRO
        self._fijar_fuente(leyenda)
        leyenda.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        leyenda.paragraph_format.space_after = Pt(10)
        leyenda.paragraph_format.line_spacing = 1.0

        # Las listas heredan de Normal para no arrastrar la fuente de tema.
        for nombre in ("List Bullet", "List Number"):
            estilo = self.doc.styles[nombre]
            estilo.base_style = normal
            estilo.font.size = Pt(TAMANO_PT)
            estilo.font.color.rgb = NEGRO
            self._fijar_fuente(estilo)

    def _fijar_fuente(self, estilo):
        """Impone Times New Roman en un estilo y elimina las referencias a las
        fuentes del tema de Word (asciiTheme, hAnsiTheme...). Sin esto Word da
        prioridad al tema y la marca de parrafo de los titulos se queda en
        Calibri, aunque el texto visible ya sea Times New Roman."""
        estilo.font.name = FUENTE
        r_pr = estilo.element.get_or_add_rPr()
        fuentes = r_pr.rFonts
        for atributo in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            clave = qn("w:%s" % atributo)
            if clave in fuentes.attrib:
                del fuentes.attrib[clave]
        for atributo in ("ascii", "hAnsi", "eastAsia", "cs"):
            fuentes.set(qn("w:%s" % atributo), FUENTE)

    # ---------------------------------------------------------------- figuras
    def fig(self, etiqueta):
        """Devuelve el numero de figura asociado a una etiqueta."""
        return str(self._figuras[etiqueta])

    # -------------------------------------------------------------- titulares
    def titulo_portada(self, texto, negrita=True, espacio_despues=6):
        """Linea centrada de la caratula."""
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(espacio_despues)
        _escribir(p, texto, negrita=negrita)
        return p

    def titulo(self, nivel, texto):
        """Titulo de seccion (1), subseccion (2) o sub-subseccion (3)."""
        p = self.doc.add_paragraph(style="Heading %d" % nivel)
        _escribir(p, texto, negrita=True)
        return p

    # --------------------------------------------------------------- parrafos
    def parrafo(self, texto="", rojo=False, centrado=False, negrita=False,
                cursiva=False, espacio_despues=6, sangria_izq=0.0):
        p = self.doc.add_paragraph()
        p.alignment = (WD_ALIGN_PARAGRAPH.CENTER if centrado
                       else WD_ALIGN_PARAGRAPH.JUSTIFY)
        p.paragraph_format.space_after = Pt(espacio_despues)
        p.paragraph_format.left_indent = Cm(sangria_izq)
        _escribir(p, texto, negrita=negrita, cursiva=cursiva, rojo=rojo)
        return p

    def espacio(self, puntos=6):
        """Linea en blanco de separacion."""
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(puntos)
        return p

    def salto_pagina(self):
        """Equivalente a \\clearpage."""
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def linea_horizontal(self):
        """Regla horizontal simple: borde inferior negro de un parrafo vacio."""
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(12)
        bordes = OxmlElement("w:pBdr")
        inferior = OxmlElement("w:bottom")
        inferior.set(qn("w:val"), "single")
        inferior.set(qn("w:sz"), "6")
        inferior.set(qn("w:space"), "1")
        inferior.set(qn("w:color"), "auto")
        bordes.append(inferior)
        p._p.get_or_add_pPr().append(bordes)
        return p

    # ----------------------------------------------------------------- listas
    def vinetas(self, elementos, rojo=False):
        """Lista con vinetas (equivalente a itemize)."""
        for elemento in elementos:
            p = self.doc.add_paragraph(style="List Bullet")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.left_indent = Cm(1.0)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            _escribir(p, elemento, rojo=rojo)

    def numerada(self, elementos, rojo=False):
        """Lista numerada que reinicia en 1 (equivalente a enumerate).

        Word continua la numeracion entre listas consecutivas que comparten
        definicion, asi que cada llamada crea su propia definicion con
        startOverride para que vuelva a empezar en 1."""
        num_id = self._nueva_numeracion()
        for elemento in elementos:
            p = self.doc.add_paragraph(style="List Number")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.left_indent = Cm(1.1)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            self._asignar_numeracion(p, num_id)
            _escribir(p, elemento, rojo=rojo)

    def _numbering(self):
        """Devuelve (creando si hace falta) la parte numbering.xml."""
        parte = self.doc.part.numbering_part
        return parte.element

    def _nueva_numeracion(self):
        """Crea una definicion <w:num> nueva que reinicia el conteo en 1 y
        devuelve su numId."""
        numbering = self._numbering()
        # numId de la lista numerada base que trae la plantilla de Word.
        abstract_id = self._abstract_id_lista_numerada(numbering)
        usados = [int(n.get(qn("w:numId")))
                  for n in numbering.findall(qn("w:num"))]
        nuevo_id = max(usados) + 1 if usados else 1

        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(nuevo_id))
        ref = OxmlElement("w:abstractNumId")
        ref.set(qn("w:val"), str(abstract_id))
        num.append(ref)
        for nivel in range(9):
            override = OxmlElement("w:lvlOverride")
            override.set(qn("w:ilvl"), str(nivel))
            inicio = OxmlElement("w:startOverride")
            inicio.set(qn("w:val"), "1")
            override.append(inicio)
            num.append(override)
        numbering.append(num)
        return nuevo_id

    def _abstract_id_lista_numerada(self, numbering):
        """abstractNumId al que apunta el estilo 'List Number' de la plantilla."""
        estilo = self.doc.styles["List Number"]
        num_pr = estilo.element.find(qn("w:pPr"))
        if num_pr is not None:
            num_id_el = num_pr.find(qn("w:numPr"))
            if num_id_el is not None:
                valor = num_id_el.find(qn("w:numId"))
                if valor is not None:
                    buscado = valor.get(qn("w:val"))
                    for num in numbering.findall(qn("w:num")):
                        if num.get(qn("w:numId")) == buscado:
                            return num.find(qn("w:abstractNumId")).get(qn("w:val"))
        # Reserva: el primer abstractNum decimal disponible.
        abstractos = numbering.findall(qn("w:abstractNum"))
        return abstractos[0].get(qn("w:abstractNumId")) if abstractos else "0"

    @staticmethod
    def _asignar_numeracion(parrafo, num_id):
        """Apunta el parrafo a la definicion de numeracion indicada."""
        p_pr = parrafo._p.get_or_add_pPr()
        antiguo = p_pr.find(qn("w:numPr"))
        if antiguo is not None:
            p_pr.remove(antiguo)
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        numero = OxmlElement("w:numId")
        numero.set(qn("w:val"), str(num_id))
        num_pr.append(ilvl)
        num_pr.append(numero)
        p_pr.append(num_pr)

    # ----------------------------------------------------------------- tablas
    def tabla(self, filas, anchos_cm=None, encabezado=True):
        """Tabla SIN FORMATO: unicamente lineas simples negras de 0.5 pt.
        Sin sombreados, sin bandas de color y sin estilos decorativos.

        filas     : lista de listas de cadenas (la primera fila es el encabezado).
        anchos_cm : anchos de columna en centimetros.
        """
        n_col = len(filas[0])
        if anchos_cm is None:
            anchos_cm = [ANCHO_UTIL_CM / n_col] * n_col

        tabla = self.doc.add_table(rows=len(filas), cols=n_col)
        tabla.style = None                     # sin estilo decorativo de Word
        tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
        tabla.autofit = False
        self._bordes_simples(tabla)

        for i, fila in enumerate(filas):
            es_encabezado = encabezado and i == 0
            if es_encabezado:
                self._repetir_encabezado(tabla.rows[i])
            for j, contenido in enumerate(fila):
                celda = tabla.cell(i, j)
                celda.width = Cm(anchos_cm[j])
                p = celda.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                _escribir(p, contenido, negrita=es_encabezado)

        # Las anchuras deben repetirse celda a celda para que Word las respete.
        for fila_tabla in tabla.rows:
            for j, celda in enumerate(fila_tabla.cells):
                celda.width = Cm(anchos_cm[j])

        self.espacio(6)
        return tabla

    @staticmethod
    def _bordes_simples(tabla):
        """Lineas negras de 0.5 pt en todo el contorno y el interior."""
        tbl_pr = tabla._tbl.tblPr
        bordes = OxmlElement("w:tblBorders")
        for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
            elemento = OxmlElement("w:%s" % lado)
            elemento.set(qn("w:val"), "single")
            elemento.set(qn("w:sz"), "4")        # 4 octavos de punto = 0.5 pt
            elemento.set(qn("w:space"), "0")
            elemento.set(qn("w:color"), "auto")  # negro automatico, sin color
            bordes.append(elemento)
        tbl_pr.append(bordes)

    @staticmethod
    def _repetir_encabezado(fila):
        """Marca la fila como encabezado para que se repita al cambiar de pagina."""
        tr_pr = fila._tr.get_or_add_trPr()
        cabecera = OxmlElement("w:tblHeader")
        cabecera.set(qn("w:val"), "true")
        tr_pr.append(cabecera)

    # ---------------------------------------------------------------- imagenes
    def imagen(self, ruta, ancho_cm=None, leyenda=None, etiqueta=None):
        """Inserta una imagen centrada EN EL PUNTO EXACTO en que se escribe
        (equivalente al especificador [H] de LaTeX: nunca flota a otra pagina),
        seguida de su leyenda numerada."""
        if not os.path.exists(ruta):
            # Si falta el archivo se deja constancia en rojo en lugar de abortar
            # la generacion del documento completo.
            self.parrafo("[Falta la imagen: %s]" % ruta, rojo=True, centrado=True)
            if leyenda is not None:
                self._leyenda(leyenda, etiqueta)
            return None

        ancho_cm = ancho_cm or ANCHO_UTIL_CM * 0.93
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True   # la leyenda no se separa
        p.add_run().add_picture(ruta, width=Cm(ancho_cm))
        if leyenda is not None:
            self._leyenda(leyenda, etiqueta)
        return p

    def _leyenda(self, texto, etiqueta):
        p = self.doc.add_paragraph(style="Caption")
        if etiqueta:
            _escribir(p, "Figura %s. " % self.fig(etiqueta), negrita=True)
        _escribir(p, texto)
        return p

    # --------------------------------------------------------------- guardado
    def guardar(self, ruta_salida):
        """Escribe el archivo .docx en disco y devuelve su ruta."""
        self.doc.save(ruta_salida)
        return ruta_salida
